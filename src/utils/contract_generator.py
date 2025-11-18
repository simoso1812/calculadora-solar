"""
Utilidad para generar contratos en formato Word.
"""
import io
import datetime
import streamlit as st
from docx import Document
from num2words import num2words

def generar_contrato_docx(datos_contrato):
    """
    Carga la plantilla de Word, reemplaza los placeholders y devuelve el documento en bytes.
    """
    try:
        # Cargamos la plantilla desde la carpeta assets
        doc = Document('assets/contrato_plantilla.docx')

        # Creamos un diccionario con los placeholders y sus valores
        context = {
            '{{NOMBRE_CLIENTE}}': datos_contrato.get('Cliente', ''),
            '{{DOCUMENTO_CLIENTE}}': datos_contrato.get('Documento del Cliente', ''),
            '{{DIRECCION_PROYECTO}}': datos_contrato.get('Dirección del Proyecto', ''),
            '{{TAMANO_DEL_SISTEMA_KWP}}': str(datos_contrato.get('Tamano del Sistema (kWp)', '')),
            '{{CANTIDAD_PANELES}}': str(datos_contrato.get('Cantidad de Paneles', '')).split(' ')[0],
            '{{POTENCIA_PANEL}}': str(datos_contrato.get('Potencia de Paneles', '')),
            '{{INVERSOR_RECOMENDADO}}': datos_contrato.get('Inversor Recomendado', ''),
            '{{VALOR_TOTAL_PROYECTO_NUMEROS}}': datos_contrato.get('Valor Total del Proyecto (COP)', ''),
            '{{FECHA_FIRMA}}': datos_contrato.get('Fecha de la Propuesta', '').strftime('%d de %B de %Y'),
        }

        # Convertimos el valor numérico a letras
        try:
            valor_str = datos_contrato.get('Valor Total del Proyecto (COP)', '$0')
            # Limpiar el string de caracteres no numéricos
            valor_limpio = valor_str.replace('$', '').replace(',', '').replace(' ', '')
            valor_numerico = int(float(valor_limpio))
            context['{{VALOR_TOTAL_PROYECTO_LETRAS}}'] = num2words(valor_numerico, lang='es').upper() + " PESOS M/CTE"
        except (ValueError, TypeError, AttributeError) as e:
            st.warning(f"No se pudo convertir el valor a letras: {e}")
            context['{{VALOR_TOTAL_PROYECTO_LETRAS}}'] = "CERO PESOS M/CTE"
        
        # Convertir fecha a español
        try:
            fecha = datos_contrato.get('Fecha de la Propuesta', datetime.date.today())
            
            # Manejar diferentes tipos de fecha
            if isinstance(fecha, str):
                # Intentar diferentes formatos de fecha
                formatos_fecha = ['%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d']
                fecha_parseada = None
                
                for formato in formatos_fecha:
                    try:
                        fecha_parseada = datetime.datetime.strptime(fecha, formato).date()
                        break
                    except ValueError:
                        continue
                
                if fecha_parseada:
                    fecha = fecha_parseada
                else:
                    raise ValueError(f"No se pudo parsear la fecha: {fecha}")
            
            # Diccionario de meses en español
            meses_espanol = {
                1: 'enero', 2: 'febrero', 3: 'marzo', 4: 'abril',
                5: 'mayo', 6: 'junio', 7: 'julio', 8: 'agosto',
                9: 'septiembre', 10: 'octubre', 11: 'noviembre', 12: 'diciembre'
            }
            
            # Formatear fecha en español
            dia = fecha.day
            mes = meses_espanol[fecha.month]
            año = fecha.year
            
            fecha_espanol = f"{dia} de {mes} de {año}"
            context['{{FECHA_FIRMA}}'] = fecha_espanol
            
        except Exception as e:
            st.warning(f"No se pudo formatear la fecha en español: {e}")
            # Fallback a formato básico
            try:
                if hasattr(datos_contrato.get('Fecha de la Propuesta', ''), 'strftime'):
                    context['{{FECHA_FIRMA}}'] = datos_contrato.get('Fecha de la Propuesta', '').strftime('%d/%m/%Y')
                else:
                    context['{{FECHA_FIRMA}}'] = "fecha no disponible"
            except:
                context['{{FECHA_FIRMA}}'] = "fecha no disponible"

        # Reemplazamos los placeholders en los párrafos
        for p in doc.paragraphs:
            for key, value in context.items():
                if key in p.text:
                    inline = p.runs
                    # Reemplazamos el texto manteniendo el formato original
                    for i in range(len(inline)):
                        if key in inline[i].text:
                            text = inline[i].text.replace(key, value)
                            inline[i].text = text
        
        # Guardamos el documento en la memoria
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        return file_stream.getvalue()

    except Exception as e:
        st.error(f"Error al generar el contrato: {e}")
        return None

